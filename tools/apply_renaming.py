"""
Applies the approved item renaming to ERPNext.

    python tools/apply_renaming.py              # dry run - shows what it would do
    python tools/apply_renaming.py --apply      # writes to ERPNext

Reads data/renaming/MRPPL-Item-Renaming.xlsx back, so whatever the Maintenance
Manager typed into the Decision and amended-name columns is what happens. The
workbook is the instruction, not the JSON it was built from.

--- What is actually changed ------------------------------------------------

Only `item_name` - the descriptive name on the ERPNext Item. The document name
(SAP1..SAP654) is the item code, and it is never touched: it is what every
stock ledger entry, bin, purchase order and shelf label points at, and renaming
a document in Frappe rewrites every one of those links. The catalog shows
`item_name`, so changing it is the whole of the visible rename and none of the
risk.

Every change is written to data/renaming/rename-log.csv before it is made, so
the previous name of every item is on disk and the whole thing can be put back.

Decisions, per row of the Renaming sheet:

    blank or "Approve"   use the proposed name
    "Reject"             leave the item alone
    "Amend"              use whatever is in "Your name, if amending"

The Optional field order sheet is the other way round: nothing happens unless a
row is explicitly marked Approve, because those items are already compliant.
"""

import argparse
import csv
import json
import os
import sys
import urllib.parse
import urllib.request

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
DATA = os.path.join(ROOT, "data", "renaming")
BOOK = os.path.join(DATA, "MRPPL-Item-Renaming.xlsx")
LOG = os.path.join(DATA, "rename-log.csv")


def erp_settings():
    """The integration credentials, read from the server's own .env."""
    env = {}
    path = os.path.join(ROOT, "server", ".env")
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, _, value = line.partition("=")
            env[key.strip()] = value.strip().strip('"').strip("'")

    url = env.get("ERPNEXT_URL", "").rstrip("/")
    key = env.get("ERPNEXT_API_KEY", "")
    secret = env.get("ERPNEXT_API_SECRET", "")
    if not (url and key and secret):
        sys.exit("ERPNEXT_URL / ERPNEXT_API_KEY / ERPNEXT_API_SECRET are not set in server/.env")
    return url, f"token {key}:{secret}"


def call(url, auth, method, path, body=None):
    request = urllib.request.Request(
        url + path,
        method=method,
        data=json.dumps(body).encode() if body else None,
        headers={"Authorization": auth, "Content-Type": "application/json"},
    )
    with urllib.request.urlopen(request) as response:
        return json.load(response)["data"]


def read_sheet(ws, columns):
    """Rows as dicts, keyed by the column names given, skipping the 4-row header."""
    out = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not row or not row[0]:
            continue
        out.append({name: (row[i] if i < len(row) else None) for i, name in enumerate(columns)})
    return out


def decide(row, proposed_key, amended_key):
    """The name this row asks for, or None to leave the item alone."""
    decision = str(row.get("decision") or "").strip().lower()
    amended = str(row.get(amended_key) or "").strip()

    if decision == "reject":
        return None
    if decision == "amend":
        return amended or None
    return str(row.get(proposed_key) or "").strip() or None


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true", help="write to ERPNext")
    args = parser.parse_args()

    if not os.path.exists(BOOK):
        sys.exit(f"{BOOK} is not there. Build it first with tools/build_rename_sheet.py")

    wb = load_workbook(BOOK, data_only=True)

    renaming = read_sheet(
        wb["Renaming"],
        ["code", "current", "proposed", "what", "why", "decision", "amended",
         "category", "unit", "qty", "rack"],
    )
    optional = read_sheet(
        wb["Optional field order"],
        ["code", "current", "proposed", "decision", "category"],
    )
    hard = read_sheet(
        wb["Needs a decision"],
        ["code", "current", "proposed", "remaining", "yours"],
    )

    planned = []

    for row in renaming:
        wanted = decide(row, "proposed", "amended")
        if wanted and wanted != row["current"]:
            planned.append((row["code"], row["current"], wanted, "Renaming"))

    # The opposite default: already compliant, so only an explicit Approve moves it.
    for row in optional:
        if str(row.get("decision") or "").strip().lower() == "approve":
            wanted = str(row.get("proposed") or "").strip()
            if wanted and wanted != row["current"]:
                planned.append((row["code"], row["current"], wanted, "Optional field order"))

    # These have no usable proposal; only a hand-written name counts.
    for row in hard:
        wanted = str(row.get("yours") or "").strip()
        if wanted and wanted != row["current"]:
            planned.append((row["code"], row["current"], wanted, "Needs a decision"))

    if not planned:
        print("Nothing to do — no row asks for a change.")
        return

    print(f"{'APPLYING' if args.apply else 'DRY RUN'}: {len(planned)} item(s) would be renamed.\n")
    for code, before, after, source in planned[:15]:
        print(f"  {code:8} {before[:44]:46} -> {after[:44]}")
    if len(planned) > 15:
        print(f"  … and {len(planned) - 15} more")

    if not args.apply:
        print("\nNothing was written. Re-run with --apply to make these changes.")
        return

    url, auth = erp_settings()

    # The rollback file is written before the first change, not after the last.
    os.makedirs(DATA, exist_ok=True)
    with open(LOG, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["item_code", "previous_name", "new_name", "sheet"])
        writer.writerows(planned)
    print(f"\nRollback written to {LOG}")

    done = 0
    failed = []
    for code, before, after, _ in planned:
        try:
            call(url, auth, "PUT",
                 "/api/resource/Item/" + urllib.parse.quote(code),
                 {"item_name": after})
            done += 1
        except Exception as error:  # noqa: BLE001 - reported, not swallowed
            failed.append((code, str(error)[:120]))

    print(f"\nRenamed {done} of {len(planned)}.")
    for code, message in failed:
        print(f"  ! {code}: {message}")


if __name__ == "__main__":
    main()
