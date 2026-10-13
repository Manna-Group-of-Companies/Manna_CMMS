"""What the 10 Oct 2026 Manna Rubber clean-up wrote to SAP, as one workbook for the store.

    python tools/build_sap_cleanup_applied.py <plan.json> <before.json> <after.json> <gr.json> <out.xlsx>

<before.json> / <after.json> are Sap-Get reads of group 108 ([0].rows), <gr.json> the Goods Receipts whose
Reference2 starts with PSC (with DocumentLines). Lists only, no formulas.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

plan_path, before_path, after_path, gr_path, out_path = sys.argv[1:6]
plan = json.loads(Path(plan_path).read_text(encoding="utf-8-sig"))
before = {str(r["ItemCode"]).strip(): r for r in json.loads(Path(before_path).read_text(encoding="utf-8-sig"))[0]["rows"]}
after = {str(r["ItemCode"]).strip(): r for r in json.loads(Path(after_path).read_text(encoding="utf-8-sig"))[0]["rows"]}
grs = json.loads(Path(gr_path).read_text(encoding="utf-8-sig"))[0]["rows"]

FONT = Font(name="Arial", size=10)
HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F4E78")
WRAP = Alignment(wrap_text=True, vertical="top")


def num(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def min_of(it):
    if it.get("ManageStockByWarehouse") == "tYES":
        for w in it.get("ItemWarehouseInfoCollection") or []:
            if str(w.get("WarehouseCode", "")).strip() == "01":
                return num(w.get("MinimalStock"))
    return num(it.get("MinInventory"))


def stamp(n):
    return "PS-" + (n["name"].upper() if n.get("combined") else str(n["physicalIds"]).split(",")[0].strip())


by_sww = {str(r.get("SWW") or "").strip(): code for code, r in after.items() if r.get("SWW")}
gr_of = {}
for d in grs:
    for l in d.get("DocumentLines") or []:
        gr_of.setdefault(str(l["ItemCode"]).strip(), []).append(str(d["DocNum"]))

wb = Workbook()


def sheet(title, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = HEAD
        c.fill = FILL
        c.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = FONT
            c.alignment = WRAP
    for i, h in enumerate(headers, 1):
        longest = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None] or [10])
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, min(max(10, longest + 2), 60))
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


# --- renamed
ren = []
for r in plan["renames"]:
    code = str(r["sapCode"]).strip()
    a = after.get(code, {})
    ren.append([code, r["currentName"], a.get("ItemName", r["newName"]), a.get("U_SubTypeA", ""), a.get("U_SubTypeB", ""),
                min_of(a) if a else r.get("minStock"), num(a.get("QuantityOnStock")) if a else None, r.get("physicalQty"),
                r.get("physicalId"), f"{r.get('band', '')} {r.get('score', '')}".strip(),
                "counted and SAP figures differ - SAP stock left as it is" if num(r.get("physicalQty")) != num(a.get("QuantityOnStock")) else ""])
sheet("Renamed", ["SAP code", "Old name", "New name", "Sub-category A", "Sub-category B", "Min stock", "SAP stock (unchanged)",
                  "Counted qty", "Counted id", "Match", "Note"], ren)

# --- new items
new = []
for n in plan["newItems"]:
    code = by_sww.get(stamp(n), "")
    a = after.get(code, {})
    notes = []
    if n.get("combined"):
        notes.append(f"one item for all counted {n['name'].lower()}s ({len(str(n['physicalIds']).split(','))} sizes); minimum = sum of the counted minimums")
    if n.get("excluded"):
        notes.append("not received: " + n["excluded"])
    if not n.get("combined") and "," in str(n.get("physicalIds", "")):
        notes.append("two counted rows with the same standard name, added up")
    if num(n.get("physicalQty")) < 0:
        notes.append("counted as a negative quantity - nothing received")
    elif num(n.get("physicalQty")) == 0:
        notes.append("counted with no stock - created without stock")
    if n.get("uomNote"):
        notes.append(n["uomNote"])
    new.append([code or "(not found by stamp)", a.get("ItemName", n["name"]), a.get("U_SubTypeA", n.get("subA", "")), a.get("U_SubTypeB", n.get("subB", "")),
                a.get("InventoryUOM", n.get("uom")), n.get("physicalQty"), num(a.get("QuantityOnStock")) if a else None, min_of(a) if a else n.get("minStock"),
                n.get("physicalIds"), ", ".join(gr_of.get(code, [])), "; ".join(notes)])
sheet("New items", ["SAP code", "Name", "Sub-category A", "Sub-category B", "Unit", "Counted qty", "Now in SAP", "Min stock",
                    "Counted id(s)", "Goods Receipt", "Note"], new, {"Counted id(s)": 40, "Note": 60})

# --- deactivated
de = []
for d in plan["deactivate"]:
    code = str(d["sapCode"]).strip()
    b = before.get(code, {})
    a = after.get(code, {})
    de.append([code, b.get("ItemName", d.get("name")), num(a.get("QuantityOnStock")) if a else None,
               "inactive" if a.get("Valid") == "tNO" and a.get("Frozen") == "tYES" else f"Valid {a.get('Valid')} Frozen {a.get('Frozen')}",
               b.get("U_SubTypeA", ""), b.get("U_SubTypeB", "")])
sheet("Deactivated", ["SAP code", "Name", "Stock", "State now", "Sub-category A", "Sub-category B"], de)

# --- left as is
man = []
for m in plan["manual"]:
    code = str(m["sapCode"]).strip()
    a = after.get(code, {})
    man.append([code, a.get("ItemName", m.get("name") or m.get("currentName")), num(a.get("QuantityOnStock")) if a else None,
                a.get("U_SubTypeA", ""), a.get("U_SubTypeB", ""), m.get("reason") or m.get("why") or m.get("note") or m.get("closest") or ""])
sheet("Left as is", ["SAP code", "Name", "Stock", "Sub-category A", "Sub-category B", "Why"], man, {"Why": 70})

# --- goods receipts
g = []
for d in grs:
    lines = d.get("DocumentLines") or []
    g.append([d["DocNum"], d["DocEntry"], str(d.get("DocDate", ""))[:10], d.get("Series"), d.get("Reference2"), len(lines),
              round(sum(num(l.get("Quantity")) for l in lines), 1), num(d.get("DocTotal")), d.get("Comments", "")])
sheet("Goods Receipts", ["Doc no.", "DocEntry", "Date", "Series", "Reference2", "Lines", "Pieces", "Doc total", "Comments"], g, {"Comments": 80})

# --- summary first
ws = wb.active
ws.title = "Summary"
active_after = [r for r in after.values() if r.get("Valid") == "tYES" and r.get("Frozen") != "tYES"]
active_before = [r for r in before.values() if r.get("Valid") == "tYES" and r.get("Frozen") != "tYES"]
rows = [
    ["Manna Rubber Products - SAP engineering items (group 108) brought in line with the physical count, 10 Oct 2026", ""],
    ["Written through the SAP Service Layer, MANNA_RUBBER_LIVE only. Every write is in the journal; every item touched was backed up first.", ""],
    ["", ""],
    ["Renamed to the standard name (exact matches), with Sub-category A/B and minimum stock", len(plan["renames"])],
    ["Created (counted items SAP lacked), at nil value", len(plan["newItems"])],
    ["   of which created with their counted stock", sum(1 for n in plan["newItems"] if num(n.get("physicalQty")) > 0)],
    ["   of which counted with no stock (created without stock)", sum(1 for n in plan["newItems"] if num(n.get("physicalQty")) <= 0)],
    ["Opening stock received (Goods Receipts, price 0, warehouse 01)", f"{len(grs)} documents, {sum(len(d.get('DocumentLines') or []) for d in grs)} lines, {round(sum(num(l.get('Quantity')) for d in grs for l in d.get('DocumentLines') or []), 1)} pcs"],
    ["Deactivated (matched nothing, zero stock)", len(plan["deactivate"])],
    ["Left as they are (matched nothing, hold stock) - incl. ENG - 399 kept after PO 1033 was cancelled", len(plan["manual"])],
    ["Group 108 before", f"{len(before)} items, {len(active_before)} active"],
    ["Group 108 after", f"{len(after)} items, {len(active_after)} active"],
    ["", ""],
    ["Notes", ""],
]
for n in plan.get("notes") or []:
    rows.append(["", n])
rows += [
    ["", "Renames change the name, sub-categories and minimum only; SAP's stock figure on those 54 items is untouched (46 differ from the count - see the Renamed sheet)."],
    ["", "Bolt / Nut / Washer: one SAP item each; SAP will warn when the total falls below the minimum, not when one size runs out."],
    ["", "New item codes from ENG-760 up are one higher than in the inspected list (one pulley was added at ENG-760 when ENG-245 was dropped from the renames)."],
    ["", "The CMMS mirror (ERPNext) follows on the first flagged sync with pull_from_sap switched on."],
]
for r in rows:
    ws.append(r)
for row in ws.iter_rows():
    for c in row:
        c.font = FONT
        c.alignment = WRAP
ws["A1"].font = Font(name="Arial", size=12, bold=True)
ws.column_dimensions["A"].width = 95
ws.column_dimensions["B"].width = 80

Path(out_path).parent.mkdir(parents=True, exist_ok=True)
try:
    wb.save(out_path)
    print("written", out_path)
except PermissionError:
    alt = str(Path(out_path).with_name(Path(out_path).stem + "-copy.xlsx"))
    wb.save(alt)
    print("written (original locked)", alt)
print("manual keys:", sorted(plan["manual"][0].keys()))
print("new items not found by stamp:", sum(1 for r in new if r[0].startswith("(")))
