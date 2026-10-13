"""
Builds the DTW import files and the review workbook for the Manna Rubber Products
SAP clean-up (10 Oct 2026), from server/scripts/planMrpplSapCleanup.js output.

    python tools/build_sap_cleanup_dtw.py plan.json <out folder>

Writes, for the user to inspect and import with DTW (oItems):
    01_Rename_Items.txt        ItemName -> the standard name, Sub-category A / B as counted
    02_Deactivate_Items.txt    Valid = tNO, Frozen = tYES
    03_New_Items.txt           new ENG- items from the physical count (one Bolt, one Nut, one Washer)
    04_Opening_Stock_GR_Header.txt / _Lines.txt
                               one Goods Receipt (oInventoryGenEntry) putting the counted
                               quantity of the new items in at nil value (price 0)
    MRPPL-SAP-Cleanup-Review.xlsx   the same rows with old and new side by side,
                                    plus the by-hand list (no DTW) and what was left out
    README.txt

The DTW files are tab-delimited with the two header rows earlier DTW imports on
this server used (both rows the oItems field names). Columns after the import
fields (CurrentName, PhysicalName, ...) are for comparing only: DTW leaves a
column it cannot map unmapped. Written in Windows-1252 (DTW's default
encoding); a character it cannot hold is replaced and reported.
"""

import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PLAN, OUTDIR = sys.argv[1], sys.argv[2]
plan = json.load(open(PLAN, encoding="utf-8"))
os.makedirs(OUTDIR, exist_ok=True)
replaced = []


def clean(v):
    if v is None:
        return ""
    s = str(v).replace("\t", " ").replace("\r", " ").replace("\n", " ")
    try:
        s.encode("cp1252")
        return s
    except UnicodeEncodeError:
        out = s.encode("cp1252", errors="replace").decode("cp1252")
        replaced.append((s, out))
        return out


def write_dtw(name, fields, info, rows):
    """fields: the oItems columns DTW imports; info: comparison-only columns after them."""
    path = os.path.join(OUTDIR, name)
    head = "\t".join(fields + info)
    with open(path, "w", encoding="cp1252", newline="") as f:
        f.write(head + "\r\n" + head + "\r\n")
        for r in rows:
            f.write("\t".join(clean(r.get(c)) for c in fields + info) + "\r\n")
    return path


# ------------------------------------------------------------------ DTW files
renames = [r for r in plan["renames"] if not r.get("unchanged")]
write_dtw(
    "01_Rename_Items.txt",
    ["ItemCode", "ItemName", "U_SubTypeA", "U_SubTypeB"],
    ["CurrentName", "CurrentSubTypeA", "CurrentSubTypeB", "PhysicalName", "Match", "Score"],
    [{"ItemCode": r["sapCode"], "ItemName": r["newName"], "U_SubTypeA": r["subA"], "U_SubTypeB": r["subB"],
      "CurrentName": r["currentName"], "CurrentSubTypeA": r["currentSubA"], "CurrentSubTypeB": r["currentSubB"],
      "PhysicalName": r["physicalName"], "Match": r["band"], "Score": r["score"]} for r in renames],
)
write_dtw(
    "02_Deactivate_Items.txt",
    ["ItemCode", "Valid", "Frozen"],
    ["ItemName", "StockInSAP", "Reason"],
    [{"ItemCode": d["sapCode"], "Valid": "tNO", "Frozen": "tYES", "ItemName": d["name"], "StockInSAP": d["sapQty"], "Reason": d["reason"]} for d in plan["deactivate"]],
)
new_fields = ["ItemCode", "ItemName", "ItemsGroupCode", "ItemType", "InventoryItem", "SalesItem", "PurchaseItem",
              "InventoryUOM", "PurchaseUnit", "SalesUnit", "ProcurementMethod", "Series", "U_SubTypeA", "U_SubTypeB", "MinInventory"]
write_dtw(
    "03_New_Items.txt",
    new_fields,
    ["PhysicalName", "PhysicalIDs", "PhysicalQty", "Location"],
    [{
        "ItemCode": n["itemCode"], "ItemName": n["name"], "ItemsGroupCode": 108, "ItemType": "itItems",
        "InventoryItem": "tYES", "SalesItem": "tYES", "PurchaseItem": "tYES",
        "InventoryUOM": n["uom"], "PurchaseUnit": n["uom"], "SalesUnit": n["uom"], "ProcurementMethod": "bom_Buy", "Series": 3,
        "U_SubTypeA": n["subA"], "U_SubTypeB": n["subB"], "MinInventory": n["minStock"] or "",
        "PhysicalName": n["physicalName"], "PhysicalIDs": n["physicalIds"], "PhysicalQty": n["physicalQty"], "Location": n["locations"],
    } for n in plan["newItems"]],
)

# Opening stock of the new items: one Goods Receipt at price 0 into the engineering store (01).
stocked = [n for n in plan["newItems"] if (n["physicalQty"] or 0) > 0]
GR_DATE = plan["generatedAt"][:10]
with open(os.path.join(OUTDIR, "04_Opening_Stock_GR_Header.txt"), "w", encoding="cp1252", newline="") as f:
    head = "RecordKey\tObjType\tDocDate\tComments"
    f.write(head + "\r\n" + head + "\r\n")
    f.write(f"0\t59\t{GR_DATE}\tPhysical stock count - opening stock of new engineering items, nil value\r\n")
write_dtw(
    "04_Opening_Stock_GR_Lines.txt",
    ["RecordKey", "LineNum", "ItemCode", "Quantity", "WarehouseCode", "UnitPrice"],
    ["ItemName"],
    [{"RecordKey": 0, "LineNum": i, "ItemCode": n["itemCode"], "Quantity": n["physicalQty"], "WarehouseCode": "01", "UnitPrice": 0, "ItemName": n["name"]} for i, n in enumerate(stocked)],
)

# ------------------------------------------------------------------ review workbook
FONT = "Arial"
HF = PatternFill("solid", fgColor="1F3A4D")
HFONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
GRID = Border(bottom=Side(style="thin", color="D9D9D9"))
NEW = PatternFill("solid", fgColor="E2EFDA")
WARN = PatternFill("solid", fgColor="FCE4D6")
INPUT = PatternFill("solid", fgColor="FFF2CC")


def sheet(wb, title, cols, rows, fills=None):
    ws = wb.create_sheet(title)
    for i, (h, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill = HFONT, HF
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    for n, r in enumerate(rows, 2):
        for i, v in enumerate(r, 1):
            c = ws.cell(row=n, column=i, value=v)
            c.font, c.border = BODY, GRID
            if fills and fills.get(i):
                f = fills[i](r)
                if f:
                    c.fill = f
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(len(rows) + 1, 2)}"
    return ws


def dropped_words(old, new):
    import re
    skip = {"NO", "NOS", "MM", "MTR", "SQ", "X", "OF", "AND", "WITH", "THE", "FOR"}
    w = lambda s: {t for t in re.split(r"[^A-Z]+", s.upper()) if len(t) > 1 and t not in skip}
    gone = sorted(w(old) - w(new))
    return ", ".join(gone)


wb = Workbook()
wb.remove(wb.active)
summary = wb.create_sheet("Summary")

sheet(wb, "1 Rename (DTW 01)",
      [("SAP code", 12), ("Current SAP name", 38), ("New name (standard)", 38), ("Words the new name drops", 22), ("Current Sub-cat A", 18), ("New Sub-cat A", 18), ("Current Sub-cat B", 16), ("New Sub-cat B", 18), ("Name as counted", 32), ("Counted ID", 10), ("Match", 9), ("Score", 7), ("Caution", 16), ("Stock in SAP", 10), ("Counted qty", 10), ("Also claimed by (now new items)", 22)],
      [[r["sapCode"], r["currentName"], r["newName"], dropped_words(r["currentName"], r["newName"]), r["currentSubA"], r["subA"], r["currentSubB"], r["subB"], r["physicalName"], r["physicalId"], r["band"], r["score"], r["caution"], r["sapQty"], r["physicalQty"], r["bumped"]] for r in renames],
      {3: lambda r: NEW, 4: lambda r: WARN if r[3] else None, 6: lambda r: NEW, 8: lambda r: NEW})
sheet(wb, "2 Deactivate (DTW 02)",
      [("SAP code", 12), ("SAP name", 42), ("Stock in SAP", 10), ("UOM", 8), ("Sub-category A", 20), ("Sub-category B", 16), ("Bolt/nut/washer", 10), ("Why", 44)],
      [[d["sapCode"], d["name"], d["sapQty"], d["uom"], d["subA"], d["subB"], "Yes" if d["fastener"] else "", d["reason"]] for d in plan["deactivate"]])
sheet(wb, "3 New items (DTW 03)",
      [("New code", 10), ("Name (standard)", 38), ("Meets convention", 10), ("What the name still needs", 30), ("Name as counted", 32), ("Counted IDs", 14), ("Counted qty", 10), ("UOM", 7), ("Unit note", 30), ("Min stock", 8), ("Sub-category A", 18), ("Sub-category B", 18), ("Brand", 12), ("Unit cost (Rs)", 10), ("Location", 14), ("Closest SAP item (not a match)", 40), ("Note", 30)],
      [[n["itemCode"], n["name"], "Yes" if n["compliant"] else "No", n["issues"], n["physicalName"], n["physicalIds"], n["physicalQty"], n["uom"], n["uomNote"], n["minStock"], n["subA"], n["subB"], n["brand"], n["unitCost"], n["locations"], n["closest"], n["bumpedFrom"]] for n in plan["newItems"]],
      {3: lambda r: WARN if r[2] == "No" else None, 7: lambda r: WARN if not r[6] else None, 9: lambda r: WARN if r[8] else None})
sheet(wb, "3b Opening stock (DTW 04)",
      [("Line", 6), ("Item code", 10), ("Name", 40), ("Quantity", 10), ("Warehouse", 10), ("Unit price", 10)],
      [[i, n["itemCode"], n["name"], n["physicalQty"], "01", 0] for i, n in enumerate(stocked)])
ws = sheet(wb, "4 Rename by hand (no DTW)",
           [("SAP code", 12), ("SAP name", 42), ("Stock in SAP", 10), ("UOM", 8), ("Sub-category A", 18), ("Bolt/nut/washer", 10), ("Open PO", 18), ("Nearest counted items", 60), ("Your new name", 36), ("Your notes", 30)],
           [[m["sapCode"], m["name"], m["sapQty"], m["uom"], m["subA"], "Yes" if m["fastener"] else "", m["openPo"], m.get("nearest", ""), None, None] for m in plan["manual"]],
           {9: lambda r: INPUT, 10: lambda r: INPUT})
sheet(wb, "5 Bolts nuts washers (count)",
      [("Counted ID", 10), ("Name as counted", 34), ("Standard name", 34), ("Category (count)", 18), ("Counted qty", 10), ("Goes into", 10)],
      [[f["id"], f["name"], f["standardName"], f["category"], f["qty"], f.get("combinedInto", "")] for f in plan["fastenersNotCreated"]])
if plan["alreadyInactive"]:
    sheet(wb, "Already inactive in SAP",
          [("SAP code", 12), ("SAP name", 42), ("Stock in SAP", 10), ("Why it was considered", 40)],
          [[d["sapCode"], d["name"], d["sapQty"], d["reason"]] for d in plan["alreadyInactive"]])

summary.column_dimensions["A"].width = 36
summary.column_dimensions["B"].width = 10
summary.column_dimensions["C"].width = 100
summary["A1"] = "Manna Rubber Products - SAP clean-up, 10 Oct 2026"
summary["A1"].font = Font(name=FONT, bold=True, size=13)
summary["A2"] = "SAP keeps the critical engineering items only, under their standard names. Built from the physical stock count and SAP read live on 10 Oct 2026."
summary["A2"].font = Font(name=FONT, italic=True, size=10, color="595959")
for c, v in zip(("A4", "B4", "C4"), ("Sheet", "Rows", "What happens")):
    summary[c] = v
    summary[c].font, summary[c].fill = HFONT, HF
lines = [
    ("1 Rename (DTW 01)", "Matched SAP items renamed to the counted item's standard name, with Sub-category A / B as the count files them. Matched = Exact or Strong, or a suspected match scoring 0.9+. Check 'Words the new name drops'."),
    ("2 Deactivate (DTW 02)", "SAP items with no match and no stock, plus matched bolts/nuts/washers with no stock: set inactive. Items on an open PO are never here."),
    ("3 New items (DTW 03)", f"Counted items with no SAP match - with or without stock - as new items ENG-{plan['maxEng'] + 1} onwards in Engineering Stocks (108); every bolt, nut and washer in one Bolt, one Nut and one Washer item. Same standard name counted twice = one item."),
    ("3b Opening stock (DTW 04)", "The counted quantity of the new items, in at nil value: one Goods Receipt, unit price 0, warehouse 01. Items counted with no stock are not on it."),
    ("4 Rename by hand (no DTW)", "SAP items with stock (or an open PO) and no match: yours to rename. Fill 'Your new name' and send it back for the DTW."),
    ("5 Bolts nuts washers (count)", "Every counted bolt, nut and washer line, and which of the three combined items its quantity goes into."),
]
for i, (s, text) in enumerate(lines):
    row = 5 + i
    summary.cell(row=row, column=1, value=s)
    summary.cell(row=row, column=2, value=f"=COUNTA('{s}'!A:A)-1")
    summary.cell(row=row, column=3, value=text)
notes = [
    "Import order: 01 Rename, 03 New items, 04 Opening stock (after 03 - it needs the items), then 02 Deactivate. Run each as a DTW Test Run first.",
    "Sub-category A / B come from the count's Main Category / Sub-Category (B left empty where it repeats A), for renamed and new items alike.",
    "New items have no HSN or tax rate yet: the first goods receipt from a bill fills them in. Their counted stock goes in at nil value (DTW 04); a later receipt's price replaces the nil value.",
    "After the imports, 'Sync with SAP' in the CMMS fetches the new names, the new items and the deactivations into Engineering Stock.",
    f"Open PO lines read 9 Oct 2026. Highest ENG code in SAP on 10 Oct: ENG-{plan['maxEng']}.",
]
for i, t in enumerate(notes):
    summary.cell(row=13 + i, column=1, value=t).font = Font(name=FONT, size=10, color="404040")
for row in summary.iter_rows(min_row=5, max_row=11):
    for c in row:
        c.font = BODY
# The user keeps the review workbook open in Excel while iterating, which locks it; a locked file
# gets a timestamped name beside it rather than failing after the DTW files are already written.
review = os.path.join(OUTDIR, "MRPPL-SAP-Cleanup-Review.xlsx")
try:
    wb.save(review)
except PermissionError:
    import datetime
    review = os.path.join(OUTDIR, f"MRPPL-SAP-Cleanup-Review-{datetime.datetime.now():%H%M}.xlsx")
    wb.save(review)
    print("review workbook was open elsewhere - written as", os.path.basename(review))

readme = f"""Manna Rubber Products - SAP clean-up DTW files (built 10 Oct 2026)

Inspect first: MRPPL-SAP-Cleanup-Review.xlsx shows every row of these files with the old and new values side by side.

Files (tab-delimited, two header rows, Windows-1252):
  01_Rename_Items.txt      {len(renames)} items   ItemCode + new ItemName
  02_Deactivate_Items.txt  {len(plan['deactivate'])} items   ItemCode + Valid=tNO + Frozen=tYES
  03_New_Items.txt         {len(plan['newItems'])} items   ENG-{plan['maxEng'] + 1}..ENG-{plan['maxEng'] + len(plan['newItems'])}, group 108, series 3
  04_Opening_Stock_GR_Header.txt + _Lines.txt
                           1 Goods Receipt, {len(stocked)} lines, unit price 0, warehouse 01

Columns after the import fields (CurrentName, PhysicalName, ItemName in 02, ...) are there to compare against.
DTW leaves a column it cannot map unmapped; if your DTW build objects, delete those columns.
If your DTW rejects tYES / tNO, replace them with Y / N.

Import (DTW -> Import -> Business Objects -> oItems, data type Tab, company MANNA_RUBBER_LIVE):
  1. 01_Rename_Items.txt     - "Update existing objects" only.  Test Run, then Import.
  2. 03_New_Items.txt        - "Add new objects" only.          Test Run, then Import.
     If SAP took an ENG- code in the meantime, renumber the ItemCode column first (and in 04).
  3. 04_Opening_Stock_GR_*   - object oInventoryGenEntry (Goods Receipt), "Add new objects".
     Header file + Lines file. Test Run, then Import. Unit price 0 = nil value; if SAP refuses a
     zero price, allow it under Document Settings or enter that line in the client.
  4. 02_Deactivate_Items.txt - "Update existing objects" only.  Test Run, then Import.
  Then press "Sync with SAP" in the CMMS (Engineering Stock).

Not in a DTW (sheet 4 of the review workbook): {len(plan['manual'])} SAP items that hold stock or sit on an open PO
and match nothing in the count. Fill "Your new name" and send it back.
"""
open(os.path.join(OUTDIR, "README.txt"), "w", encoding="utf-8").write(readme)

print("wrote", OUTDIR, f"- rename {len(renames)}, deactivate {len(plan['deactivate'])}, new {len(plan['newItems'])}, by hand {len(plan['manual'])}")
if replaced:
    print("characters Windows-1252 cannot hold were replaced:", replaced[:10])
